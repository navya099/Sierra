"""v12.7.1 KDI unit-price catalogue; standalone, no v12.6.5 engine changes."""
from __future__ import annotations
from pathlib import Path
import json, math
BASE=Path(__file__).resolve().parent

def load_catalog(path=None):
    return json.loads(Path(path or BASE/'kdi_general_rail_2019.json').read_text(encoding='utf-8'))

def load_config(path=None):
    p=Path(path or BASE/'unit_price_config.json')
    return json.loads(p.read_text(encoding='utf-8'))

def save_config(config,path=None):
    p=Path(path or BASE/'unit_price_config.json')
    p.write_text(json.dumps(config,ensure_ascii=False,indent=2),encoding='utf-8')

def effective_price(item,config):
    base=item.get('source_price_krw')
    if base is None:return None
    override=config.get('items',{}).get(item['code'],{})
    mode=override.get('mode','factor')
    if mode=='manual':
        value=override.get('unit_price')
        if value is None:raise ValueError(f"{item['code']}: 직접 입력 단가가 없습니다.")
    elif mode=='factor':
        value=base*float(override.get('factor',config.get('default_factor',1.0)))
    else:raise ValueError(f"{item['code']}: 잘못된 모드 {mode}")
    if not math.isfinite(float(value)) or float(value)<0:raise ValueError('단가는 0 이상의 유한수여야 합니다.')
    return float(value)

def validate_quantity_price(quantity,quantity_unit,item,config,used_groups=None):
    if quantity is None:raise ValueError('미산출 수량은 계산할 수 없습니다.')
    if quantity_unit!=item['unit']:raise ValueError(f"단위 불일치: {quantity_unit} / {item['unit']}")
    if item['method'] not in ('ROUGH','DETAIL'):raise ValueError('미검증 단가 또는 미매핑 항목')
    key=item['group']
    if used_groups is not None and key in used_groups and used_groups[key]!=item['method']:
        raise ValueError(f'{key}: 개략·세부 단가 중복 적용 금지')
    p=effective_price(item,config)
    if p is None:raise ValueError('단가 미검증')
    if used_groups is not None:used_groups[key]=item['method']
    return float(quantity)*p

def export_unit_prices_xlsx(path,catalog=None,config=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    catalog=catalog or load_catalog();config=config or load_config()
    wb=Workbook();wb.remove(wb.active)
    def sheet(name,header,records):
        ws=wb.create_sheet(name);ws.append(header)
        for row in records:ws.append(row)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D');c.alignment=Alignment(wrap_text=True)
        ws.row_dimensions[1].height=32
        for i in range(1,len(header)+1):ws.column_dimensions[get_column_letter(i)].width=23 if i>1 else 24
        return ws
    items=catalog['items']; overrides=config.get('items',{})
    sheet('단가기준정보',['항목','값'],[['자료명',catalog['source']],['적용분야','일반철도'],['설계기준연도',config['design_year']],['원문단가기준연도',2019],['기본보정계수',config.get('default_factor',1)],['주의','보정계수 1.0은 미보정. 미래 물가 예측값 아님'],['참고','원문 단가: 2019년 / 표 Ⅲ-66 / 보고서 p.102 / PDF p.129']])
    sheet('단가마스터',['코드','공종','구분','항목','규격','적용조건','단위','원문단가','원문단위','기준연도','원화환산단가','적용단가(원)','표번호','보고서쪽','PDF쪽','상태'],[[x['code'],x['group'],x['method'],x['name'],x['spec'],x['condition'],x['unit'],x['source_price'],x['source_scale'],x['source_year'],x['source_price_krw'],effective_price(x,config),x['source_table'],x['source_page'],x['pdf_page'],'확인' if x['source_price_krw'] is not None else '미검증'] for x in items])
    sheet('단가산출근거',['코드','산식','출처','주의'],[[x['code'],'원문 천원 × 1,000 × 보정계수 (또는 직접입력)' if x['source_price'] is not None else '미산출',f"표 {x['source_table']} / p.{x['source_page']}",x['notes']] for x in items])
    sheet('단가보정',['코드','적용모드','개별보정계수','직접단가(원)','설계연도','비고'],[[x['code'],overrides.get(x['code'],{}).get('mode','factor'),overrides.get(x['code'],{}).get('factor',config.get('default_factor',1.0)),overrides.get(x['code'],{}).get('unit_price'),config['design_year'],'단가 관리 UI 또는 JSON에서 수정'] for x in items])
    sheet('수량연계표',['수량항목','수량단위','단가코드','상태','이유'],[['절토 총량','m³','','미매핑','토사/풍화암/발파암 구성비 필요'],['성토 총량','m³','','미매핑','상부/하부노반 구성비 필요'],['교량 연장','m','','미매핑','원문 교량 형식·기초 조건 및 단위 검토 필요'],['터널 연장','m','','미매핑','원문 터널 조건·단위 검토 필요']])
    sheet('미매핑검토',['코드','공종','항목','사유'],[[x['code'],x['group'],x['name'],'원문 가격 미검증 또는 세부 수량 부재'] for x in items if x['source_price_krw'] is None])
    wb.active=0;wb.save(str(path))
