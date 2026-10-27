"""Bridge quantity takeoff v12.6.4: no invented dimensions or volumes."""
from dataclasses import dataclass
import math
import ezdxf
from .bridge_section import load_bridge_section
from .bridge_girder import span_ranges

@dataclass(frozen=True)
class BridgeQuantity:
    name: str
    start_sta: float
    end_sta: float
    length_m: float
    spans: int
    abutments: int
    piers: int
    girder_lines: int | None
    slab_area_m2: float | None
    girder_area_m2: float | None
    slab_volume_m3: float | None
    girder_volume_m3: float | None
    cap_volume_m3: float | None
    footing_volume_m3: float | None
    source: str

def _area(ent):
    pts=list(ent.get_points('xyb'))
    if not ent.closed or len(pts)<3: raise ValueError('교량 단면은 폐합 LWPOLYLINE이어야 합니다.')
    if any(abs(p[2])>1e-12 for p in pts): raise ValueError('교량 단면 bulge 원호는 지원하지 않습니다.')
    return abs(sum(x*pts[(i+1)%len(pts)][1]-pts[(i+1)%len(pts)][0]*y for i,(x,y,_) in enumerate(pts)))/2

def read_bridge_areas(path):
    # Reuse the same cross-section geometry validator as the plan drawing.
    section=load_bridge_section(path)
    doc=ezdxf.readfile(str(path))
    slab=[e for e in doc.modelspace() if e.dxf.layer.upper()=='BRIDGE_SLAB']
    girders=[e for e in doc.modelspace() if e.dxf.layer.upper()=='BRIDGE_GIRDER']
    if len(slab)!=1 or len(girders)!=len(section.girders): raise ValueError('교량 단면 레이어의 부재 개수가 일치하지 않습니다.')
    a=_area(slab[0]); b=sum(_area(e) for e in girders)
    if not math.isfinite(a+b) or a<=0 or b<=0:raise ValueError('슬래브/거더 면적이 유효하지 않습니다.')
    return a,b,len(girders)

def calculate_bridge_quantities(segments, girder_settings, plan_pier_settings, pier_settings, section_path=None):
    girder_settings.validate();plan_pier_settings.validate();pier_settings.validate()
    areas=read_bridge_areas(section_path) if section_path else None
    result=[]
    for seg in segments:
        if seg.kind!='BRIDGE':continue
        a,b=seg.start_sta,seg.end_sta
        if not all(map(math.isfinite,(a,b))) or b<=a:raise ValueError(f'{seg.name}: 교량 측점이 유효하지 않습니다.')
        n=len(span_ranges(a,b,girder_settings.girder_length)); p=n-1; L=b-a
        # Cap and footing are idealized rectangular prisms, not actual 3D solids.
        cap=p*plan_pier_settings.coping_width*plan_pier_settings.coping_length*pier_settings.cap_height
        footing=p*plan_pier_settings.foundation_width*plan_pier_settings.foundation_length*pier_settings.footing_height
        result.append(BridgeQuantity(seg.name,a,b,L,n,2,p,
            n*areas[2] if areas else None,
            areas[0] if areas else None,areas[1] if areas else None,
            areas[0]*L if areas else None,areas[1]*L if areas else None,
            cap,footing,seg.source))
    return result

def export_bridge_xlsx(rows,path,section_path=None):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill,Font
    from openpyxl.utils import get_column_letter
    wb=Workbook();ws=wb.active;ws.title='교량별 수량'
    ws.append(['교량명','시점 STA(m)','종점 STA(m)','연장(m)','경간(EA)','교대(EA)','교각(EA)',
               '거더 열·경간 합(EA)','슬래브 단면적(m²)','거더 합계 단면적(m²)',
               '슬래브 체적(m³)','거더 체적(m³)','코핑 개략체적(m³)','기초 개략체적(m³)','출처'])
    for r in rows:ws.append([r.name,r.start_sta,r.end_sta,r.length_m,r.spans,r.abutments,r.piers,
        r.girder_lines,r.slab_area_m2,r.girder_area_m2,r.slab_volume_m3,r.girder_volume_m3,
        r.cap_volume_m3,r.footing_volume_m3,r.source])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for c in row[1:14]:c.number_format='#,##0.000'
    for i in range(1,16):ws.column_dimensions[get_column_letter(i)].width=23
    sm=wb.create_sheet('집계 및 산출기준');sm.append(['항목','수량','단위'])
    fields=[('교량 개소',len(rows),'EA'),('총연장',sum(r.length_m for r in rows),'m'),
        ('경간',sum(r.spans for r in rows),'EA'),('교대',sum(r.abutments for r in rows),'EA'),
        ('교각',sum(r.piers for r in rows),'EA'),
        ('거더 열·경간 합',sum(r.girder_lines for r in rows) if rows and all(r.girder_lines is not None for r in rows) else None,'EA'),
        ('슬래브 체적',sum(r.slab_volume_m3 for r in rows) if rows and all(r.slab_volume_m3 is not None for r in rows) else None,'m³'),
        ('거더 체적',sum(r.girder_volume_m3 for r in rows) if rows and all(r.girder_volume_m3 is not None for r in rows) else None,'m³'),
        ('코핑 개략체적',sum(r.cap_volume_m3 for r in rows),'m³'),
        ('기초 개략체적',sum(r.footing_volume_m3 for r in rows),'m³')]
    for x in fields:sm.append(list(x))
    for x in [('교량 단면 DXF',str(section_path) if section_path else '미지정 (슬래브·거더 물량 공란)',''),
      ('슬래브','폐합 단면적 × 전체 교량 연장; 단면 DXF 단위 m',''),
      ('거더','각 거더 단면적 합계 × 전체 교량 연장 (거더 열 수 반영)',''),
      ('교각 코핑','교각 수 × 평면 코핑 폭 × 길이 × 종단 코핑 높이 (직육면체 가정)',''),
      ('교각 기초','교각 수 × 평면 기초 폭 × 길이 × 종단 기초 높이 (직육면체 가정)',''),
      ('미산출','교각 기둥, 교대 콘크리트, 철근, 받침, 신축이음, 접속부',''),
      ('주의','교량 단면의 슬래브·거더가 중첩되지 않는 별도 부재라는 가정; 시공 물량 아님','')]:sm.append(list(x))
    for c,w in [('A',28),('B',95),('C',15)]:sm.column_dimensions[c].width=w
    for sheet in (ws,sm):
        for c in sheet[1]:c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(bold=True,color='FFFFFF')
    wb.save(str(path))
