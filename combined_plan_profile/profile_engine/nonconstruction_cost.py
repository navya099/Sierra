"""v12.7.4-T1: independent, manual non-construction cost ledger.

Amounts are KRW, explicitly entered, never derived from construction cost.
No statutory ratios, taxes, or contingency formulas are assumed.
"""
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from collections import defaultdict
from pathlib import Path
import json

CATEGORIES = ('용지비', '부대비', '예비비', '기타 사업비')
DEFAULT_ITEMS = (
    ('용지비', '토지 매입비'),
    ('용지비', '지장물 보상비'),
    ('부대비', '조사·설계비'),
    ('부대비', '감리비'),
    ('부대비', '기타 부대비'),
    ('예비비', '예비비'),
)
SCHEMA = 'railway-nonconstruction-t1'

@dataclass(frozen=True)
class CostItem:
    category: str
    name: str
    amount: Decimal
    basis: str = ''
    note: str = ''

def parse_amount(value):
    try:
        amount = Decimal(str(value).strip().replace(',', ''))
    except (InvalidOperation, ValueError):
        raise ValueError('금액은 숫자로 입력해야 합니다.') from None
    if not amount.is_finite() or amount < 0:
        raise ValueError('금액은 0 이상의 유한한 숫자여야 합니다.')
    if amount != amount.to_integral_value():
        raise ValueError('금액은 원 단위 정수로 입력하세요.')
    return amount

def validate(items):
    result = []
    seen = set()
    for index, item in enumerate(items, 1):
        if item.category not in CATEGORIES:
            raise ValueError(f'{index}행: 지원하지 않는 비용 분류')
        name = item.name.strip()
        if not name:
            raise ValueError(f'{index}행: 항목명 누락')
        key = (item.category, name.casefold())
        if key in seen:
            raise ValueError(f'{index}행: 동일 분류에 중복 항목명 ({name})')
        seen.add(key)
        amount = parse_amount(item.amount)
        result.append(CostItem(item.category, name, amount, item.basis.strip(), item.note.strip()))
    return result

def summarize(items):
    validated = validate(items)
    sums = defaultdict(lambda: Decimal(0))
    for item in validated:
        sums[item.category] += item.amount
    return {category: sums[category] for category in CATEGORIES}

def save_json(path, items):
    items = validate(items)
    data = {'schema': SCHEMA, 'version': 1, 'currency': 'KRW',
            'items': [{'category': i.category, 'name': i.name,
                       'amount': str(i.amount), 'basis': i.basis, 'note': i.note}
                      for i in items]}
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')

def load_json(path):
    data = json.loads(Path(path).read_text(encoding='utf-8'))
    if data.get('schema') != SCHEMA or data.get('version') != 1 or data.get('currency') != 'KRW':
        raise ValueError('T1 비용 원장 형식이 아닙니다.')
    if not isinstance(data.get('items'), list):
        raise ValueError('비용 항목 목록이 없습니다.')
    return validate([CostItem(row['category'], row['name'], parse_amount(row['amount']),
                              row.get('basis', ''), row.get('note', ''))
                     for row in data['items']])

def export_xlsx(path, items):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    items = validate(items)
    totals = summarize(items)
    wb = Workbook()
    summary = wb.active
    summary.title = '공사비외비용총괄'
    summary.append(['분류', '금액(원)', '항목 수', '산정 방식'])
    for category in CATEGORIES:
        subset = [i for i in items if i.category == category]
        summary.append([category, int(totals[category]), len(subset), '사용자 직접 입력'])
    summary.append(['산정분 합계', int(sum(totals.values())), len(items), '총사업비 아님'])
    detail = wb.create_sheet('공사비외비용내역')
    detail.append(['분류', '항목명', '금액(원)', '산정 근거', '비고'])
    for item in items:
        detail.append([item.category, item.name, int(item.amount), item.basis, item.note])
    basis = wb.create_sheet('산정기준')
    for row in [
        ('항목', '내용'),
        ('범위', '용지비·부대비·예비비·기타 사업비 직접 입력분'),
        ('산정 방식', '사용자가 입력한 원 단위 금액을 단순 합산'),
        ('공사비', 'D5 공사비와 별개로 산정하며 공사비를 포함하지 않음'),
        ('세금', '부가가치세 포함 여부는 사용자 입력 근거에 명시; 자동 산정하지 않음'),
        ('예비비', '예비비율·대상 금액 자동 계산하지 않음'),
        ('중복', '같은 분류 내 동일 항목명 중복을 차단; 다른 이름/분류의 의미상 중복은 사용자가 검토'),
        ('기준연도', '가격 기준연도 및 물가 보정은 자동 적용하지 않음'),
        ('합계', '산정분 합계이며 확정 총사업비가 아님'),
    ]:
        basis.append(row)
    for ws in wb:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.alignment = Alignment(horizontal='center')
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 25 if col <= 3 else 52
    for ws, col in ((summary, 2), (detail, 3)):
        for row in ws.iter_rows(min_row=2):
            row[col-1].number_format = '#,##0"원"'
    wb.save(path)
