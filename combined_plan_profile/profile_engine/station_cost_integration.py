"""U12 station-only reconciliation. No mutation of mainline or project ledger."""
from decimal import Decimal
from .station81_rough import estimate


def reconcile(u9_counts, u11_entries, track_type, source='U11', include=False):
    if source not in ('U9', 'U11', 'NONE'):
        raise ValueError('정거장 산정 출처는 U9, U11, NONE 중 하나여야 합니다.')
    u9_records = estimate(u9_counts, track_type)
    u9_amount = sum((r['amount'] for r in u9_records), Decimal(0))
    u11_amount = sum((Decimal(row['price_krw']) for _, _, row in u11_entries if row), Decimal(0))
    u9_qty = sum((int(r['quantity']) for r in u9_records), 0)
    u11_qty = sum(1 for _, _, row in u11_entries if row)
    unresolved = [station['name'] for station, _, row in u11_entries if row is None]
    overlap_risk = bool(u9_qty and u11_qty)
    if include and source == 'U11' and unresolved:
        raise ValueError('미산정 정거장이 존재하므로 U11 전체 금액의 사업비 반영을 차단합니다: ' + ', '.join(unresolved))
    if include and source == 'NONE':
        raise ValueError('사업비 반영 시 U9 또는 U11을 선택하세요.')
    chosen = u9_amount if source == 'U9' else u11_amount if source == 'U11' else Decimal(0)
    return dict(u9_amount=u9_amount, u11_amount=u11_amount, u9_qty=u9_qty,
                u11_qty=u11_qty, unresolved=unresolved, overlap_risk=overlap_risk,
                source=source, include=include, selected_amount=chosen,
                project_addition=chosen if include else Decimal(0))


def add_sheet(wb, result):
    from openpyxl.styles import Font, PatternFill
    ws=wb.create_sheet('정거장공사비통합')
    ws.append(['항목','수량(개소)','금액(원)','상태 / 비고'])
    ws.append(['U9 수동 개소수 산정',result['u9_qty'],int(result['u9_amount']),'기존 U9 수량 기준'])
    ws.append(['U11 구간별 확정 산정',result['u11_qty'],int(result['u11_amount']),'구조물 판별 후 선택된 단가'])
    ws.append(['미산정 정거장',len(result['unresolved']),None,', '.join(result['unresolved']) or '없음'])
    ws.append(['중복 계상 가능성',None,None,'주의: U9/U11 양쪽에 산정 내역 존재' if result['overlap_risk'] else '중복 산정 내역 없음 (동일 정거장 여부는 별도 확인)'])
    ws.append(['적용 산정 방식',None,None,{'U9':'U9만 사용','U11':'U11만 사용','NONE':'반영 안 함'}[result['source']]])
    ws.append(['선택 산정금액',None,int(result['selected_amount']),'U9+U11 합산 금지'])
    ws.append(['총사업비 반영 선택',None,None,'반영 요청' if result['include'] else '미반영'])
    ws.append(['총사업비 추가 반영 후보',None,int(result['project_addition']),'참고용 별도 금액: 기존 본선/총사업비 원장 미수정'])
    ws.append(['본선 공사비 차감',None,0,'자동 차감 없음. KDI 원가 중복 범위 별도 검토'])
    ws.freeze_panes='A2'
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
    for col,width in {'A':30,'B':20,'C':26,'D':100}.items():ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=2):row[2].number_format='#,##0'
    return ws
