"""D5: merge existing KDI rough/detail XLSX by exclusive work group.

Source workbooks are treated as audited snapshots; no hidden quantity or price inference.
"""
from dataclasses import dataclass
from math import isfinite
from pathlib import Path

GROUPS=('토공','교량','터널')
@dataclass(frozen=True)
class SourceLine:
    group:str
    method:str
    code:str
    description:str
    quantity:float
    unit:str
    unit_price:float
    amount:float
    basis:str
    source_file:str

def read_cost_workbook(path, method):
    from openpyxl import load_workbook
    path=Path(path)
    if method not in ('ROUGH','DETAIL'): raise ValueError('산정 구분 오류')
    tab='공사비상세' if method=='ROUGH' else '세부공사비내역'
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        if tab not in wb.sheetnames: raise ValueError(f'{path.name}: {tab} 시트가 없습니다.')
        ws=wb[tab]
        expected=['공종','단가코드'] if method=='ROUGH' else ['공종','단가코드']
        header=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
        if header[:2]!=expected:raise ValueError(f'{path.name}: 예상하지 못한 열 구성')
        lines=[]
        for rownum,cells in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
            if not cells or not cells[0]:continue
            if method=='ROUGH':
                group,code,desc,q,unit,price,amount,basis=(*cells[:8],)
            else:
                group,code,desc,q,unit,price,amount,basis=(*cells[:8],)
            if group not in GROUPS:raise ValueError(f'{path.name} {rownum}행: 알 수 없는 공종')
            if not code:raise ValueError(f'{path.name} {rownum}행: 단가코드 누락')
            nums=[float(x) for x in (q,price,amount)]
            if not all(isfinite(x) for x in nums) or nums[0]<=0 or nums[1]<0 or nums[2]<0:
                raise ValueError(f'{path.name} {rownum}행: 수량/단가/금액 오류')
            if abs(nums[0]*nums[1]-nums[2])>max(.02,abs(nums[2])*1e-9):
                raise ValueError(f'{path.name} {rownum}행: 수량×단가와 금액 불일치')
            if (method=='ROUGH' and unit!='km') or (method=='DETAIL' and unit=='km'):
                raise ValueError(f'{path.name} {rownum}행: 산정 방식과 단위 불일치')
            lines.append(SourceLine(group,method,str(code),str(desc or ''),nums[0],str(unit),nums[1],nums[2],str(basis or ''),path.name))
        if not lines:raise ValueError(f'{path.name}: 산정 내역이 없습니다.')
        return lines
    finally:wb.close()

def integrate_costs(rough_lines,detail_lines,selected):
    """selected: {group: 'ROUGH'|'DETAIL'|'NONE'}; never add both for one group."""
    combined=list(rough_lines)+list(detail_lines)
    for line in combined:
        if line.group not in GROUPS or line.method not in ('ROUGH','DETAIL'):
            raise ValueError('유효하지 않은 산정 내역')
    chosen=[];audit=[]
    for group in GROUPS:
        mode=selected.get(group,'NONE')
        if mode not in ('ROUGH','DETAIL','NONE'):raise ValueError(f'{group}: 잘못된 산정방식')
        available={line.method for line in combined if line.group==group}
        if mode!='NONE' and mode not in available:raise ValueError(f'{group}: 선택한 {mode} 내역이 없습니다.')
        chosen.extend(line for line in combined if line.group==group and line.method==mode)
        audit.append((group,mode,' / '.join(sorted(available)) if available else '없음',
                      '미산정' if mode=='NONE' else '선택 방식만 반영 (다른 방식 제외)'))
    return chosen,audit

def export_integrated_cost(path,lines,audit,rough_file,detail_file):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    from openpyxl.utils import get_column_letter
    wb=Workbook();wb.remove(wb.active)
    def sheet(name,head,data):
        ws=wb.create_sheet(name);ws.append(head)
        for r in data:ws.append(list(r))
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
        for i in range(1,len(head)+1):ws.column_dimensions[get_column_letter(i)].width=27 if i!=4 else 18
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value,(int,float)):c.number_format='#,##0.00'
    sheet('통합공사비총괄',['공종','선택방식','산정금액(원)','상태'],
          [(g,next((a[1] for a in audit if a[0]==g),'NONE'),sum(r.amount for r in lines if r.group==g),
            '산정' if any(r.group==g for r in lines) else '미산정') for g in GROUPS]+
          [('산정분 합계','',sum(r.amount for r in lines),'총사업비 아님 / 미산정 항목 제외')])
    sheet('통합공사비내역',['공종','산정방식','단가코드','적용조건','수량','단위','적용단가(원)','금액(원)','단가출처','원본파일'],
          [(r.group,r.method,r.code,r.description,r.quantity,r.unit,r.unit_price,r.amount,r.basis,r.source_file) for r in lines])
    sheet('중복방지검토',['공종','채택방식','원본 존재방식','처리'],audit)
    sheet('산정기준',['항목','내용'],[
        ('개략 원본',str(rough_file or '미선택')),('세부 원본',str(detail_file or '미선택')),
        ('배타적 적용','공종별 개략/세부 중 정확히 하나만 채택; 나머지는 합계에서 제외'),
        ('세부 범위','세부 산정서에 입력된 항목만 포함; 세부 미산정 항목을 개략 단가로 자동 보충하지 않음'),
        ('가격','원본 XLSX의 적용단가와 금액 그대로 사용; 가격 기준연도 일치 여부 사용자 확인'),
        ('주의','동일 공종 내 서로 다른 시설/범위의 세부+개략 혼합은 지원하지 않음'),
        ('합계','산정분 합계이며 전체 공사비 또는 총사업비가 아님')])
    wb.active=0;wb.save(path)
