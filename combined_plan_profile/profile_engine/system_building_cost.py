"""v12.8-S1 independent quantity × unit price construction-cost ledger.
No KDI prices are assumed. Explicitly unestimated groups remain missing.
"""
from dataclasses import dataclass
from decimal import Decimal,InvalidOperation,ROUND_HALF_UP
from pathlib import Path
import json
import re

GROUPS=('궤도','전력','전차선','신호','통신','건축','기계설비')
SCHEMA='railway-system-building-s1'
@dataclass(frozen=True)
class Item:
    group:str
    name:str
    quantity:Decimal
    unit:str
    unit_price:Decimal
    basis:str=''
    note:str=''

def number(value, label, positive=False):
    try:d=Decimal(str(value).strip().replace(',',''))
    except (InvalidOperation,ValueError,TypeError):raise ValueError(f'{label}: 숫자를 입력하세요.') from None
    if not d.is_finite() or d<0 or (positive and d<=0):raise ValueError(f'{label}: 유효한 양수 값을 입력하세요.' if positive else f'{label}: 0 이상이어야 합니다.')
    return d

def amount(item):
    return (item.quantity*item.unit_price).quantize(Decimal('1'),rounding=ROUND_HALF_UP)

def validate(items):
    out=[];seen=set()
    for n,item in enumerate(items,1):
        if item.group not in GROUPS:raise ValueError(f'{n}행: 잘못된 공종')
        name=item.name.strip();unit=item.unit.strip()
        if not name or not unit:raise ValueError(f'{n}행: 항목명·단위는 필수입니다.')
        # KDI items are identified by the original catalog code, not by the
        # human-readable name: EARTH / BRIDGE / TUNNEL share a source name.
        match=re.search(r'KDI 단가 DB 선택:\s*([A-Za-z0-9_]+)',item.basis)
        if match:
            key=(item.group,'kdi',match.group(1).upper())
            if key in seen:raise ValueError(f'{n}행: 동일 KDI 단가 코드 중복: {match.group(1)}')
        else:
            key=(item.group,'manual',name.casefold())
            if key in seen:raise ValueError(f'{n}행: 공종 내 항목명 중복: {name}')
        seen.add(key)
        q=number(item.quantity,'수량',positive=True)
        price=number(item.unit_price,'단가')
        if price!=price.to_integral_value():raise ValueError(f'{n}행: 단가는 원 단위 정수로 입력하세요.')
        if not item.basis.strip():raise ValueError(f'{n}행: 단가·수량 산정 근거를 입력하세요.')
        out.append(Item(item.group,name,q,unit,price,item.basis.strip(),item.note.strip()))
    return out

def summary(items):
    rows=validate(items)
    return {g:(sum((amount(i) for i in rows if i.group==g),Decimal(0)) if any(i.group==g for i in rows) else None) for g in GROUPS}

def save_json(path,items):
    rows=validate(items)
    data={'schema':SCHEMA,'version':1,'currency':'KRW','items':[
        dict(group=i.group,name=i.name,quantity=str(i.quantity),unit=i.unit,unit_price=str(i.unit_price),basis=i.basis,note=i.note) for i in rows]}
    Path(path).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')

def load_json(path):
    data=json.loads(Path(path).read_text(encoding='utf-8'))
    if (data.get('schema'),data.get('version'),data.get('currency'))!=(SCHEMA,1,'KRW'):raise ValueError('S1 비용 원장 형식이 아닙니다.')
    if not isinstance(data.get('items'),list):raise ValueError('원장 항목 누락')
    return validate([Item(r['group'],r['name'],number(r['quantity'],'수량'),r['unit'],number(r['unit_price'],'단가'),r.get('basis',''),r.get('note','')) for r in data['items']])

def export_xlsx(path,items):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    from openpyxl.utils import get_column_letter
    rows=validate(items);totals=summary(rows)
    wb=Workbook();ws=wb.active;ws.title='시스템건축공사비총괄'
    ws.append(['공종','산정금액(원)','상태','항목수'])
    for g,v in totals.items():ws.append([g,int(v) if v is not None else None,'산정' if v is not None else '미산정',sum(i.group==g for i in rows)])
    ws.append(['산정분 합계',int(sum((v for v in totals.values() if v is not None),Decimal(0))),'전체 공사비 아님',len(rows)])
    ws=wb.create_sheet('시스템건축공사비내역');ws.append(['공종','항목명','수량','단위','단가(원)','금액(원)','산정 근거','비고'])
    for i in rows:ws.append([i.group,i.name,float(i.quantity),i.unit,int(i.unit_price),int(amount(i)),i.basis,i.note])
    ws=wb.create_sheet('산정기준')
    for row in [('항목','내용'),('산정방식','수량 × 사용자 입력 단가, 원 단위 반올림'),('기준단가','KDI 단가 자동 적용하지 않음; 사용자가 근거 입력'),('미산정','항목이 없는 공종은 미산정이며 0원이 아님'),('중복방지','노반 토공·교량·터널 비용은 입력 대상이 아님'),('합산','D5/P1과의 통합은 후속 단계에서 수행'),('가격 기준연도','사용자 산정 근거에 명시, 자동 물가 보정 없음'),('부가가치세','포함 여부 자동 판단하지 않음')]:ws.append(row)
    for ws in wb:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17365D')
        for col in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=18 if col<7 else 48
    for row in wb['시스템건축공사비총괄'].iter_rows(min_row=2):row[1].number_format='#,##0"원"'
    for row in wb['시스템건축공사비내역'].iter_rows(min_row=2):
        for cell in (row[4],row[5]):cell.number_format='#,##0"원"'
    wb.save(path)
