"""v12.6.2 approximate earthwork, horizontal existing ground assumption.

Cross section vertices are in local (offset, elevation) coordinates and may be
reused by a future cross-section DXF renderer. No drawing side effects.
"""
from dataclasses import dataclass
import math
from .profile_calculator import PiecewiseLinear

@dataclass(frozen=True)
class SimpleSection:
    sta: float
    ground_el: float
    plan_el: float
    height: float
    kind: str
    area_m2: float
    # Offset/elevation closed polygon; empty in neutral case
    polygon: tuple[tuple[float,float], ...]
    excluded: bool = False
    structure_type: str = ""  # 교량/터널, 경계 포함; 중첩 시 교량/터널

@dataclass(frozen=True)
class EarthworkInterval:
    start_sta: float
    end_sta: float
    cut_m3: float
    fill_m3: float
    excluded: bool

@dataclass(frozen=True)
class EarthworkResult:
    sections: tuple[SimpleSection,...]
    intervals: tuple[EarthworkInterval,...]
    cut_m3: float
    fill_m3: float
    excluded_length_m: float
    roadbed_width: float
    cut_slope: float
    fill_slope: float


def section_at(sta, ground_el, plan_el, width, cut_slope, fill_slope, excluded=False, structure_type=""):
    """Section geometry, with both side slopes horizontal:vertical = m:1."""
    h=plan_el-ground_el
    kind='FILL' if h>1e-9 else 'CUT' if h < -1e-9 else 'ZERO'
    ah=abs(h)
    m=fill_slope if h>0 else cut_slope
    area=width*ah+m*ah*ah
    half=width/2
    toe=half+m*ah
    if kind=='FILL':
        polygon=((-toe,ground_el),(-half,plan_el),(half,plan_el),(toe,ground_el))
    elif kind=='CUT':
        polygon=((-toe,ground_el),(-half,plan_el),(half,plan_el),(toe,ground_el))
    else:
        polygon=()
    return SimpleSection(sta,ground_el,plan_el,h,kind,area,polygon,excluded,structure_type)


def calculate_earthwork(ground_points, alignment, structures=(), width=10., cut_slope=1.5,
                        fill_slope=1.5, interval=25.):
    for name,value,zero_ok in [('노반 폭',width,False),('절토 사면',cut_slope,True),
                               ('성토 사면',fill_slope,True),('측점 간격',interval,False)]:
        if not math.isfinite(value) or value<0 or (not zero_ok and value==0):
            raise ValueError(f'{name}은 유효한 양수여야 합니다.')
    ground=PiecewiseLinear(list(ground_points))
    start=max(ground.xs[0],alignment.vips[0][0])
    end=min(ground.xs[-1],alignment.vips[-1][0])
    if end<=start:raise ValueError('지반고와 계획고의 공통 구간이 없습니다.')
    exclusions=[]
    typed_exclusions=[]
    for s in structures:
        if s.kind not in ('BRIDGE','TUNNEL'):continue
        a,b=max(start,s.start_sta),min(end,s.end_sta)
        if a<b:
            exclusions.append((a,b))
            typed_exclusions.append((a,b,"교량" if s.kind=="BRIDGE" else "터널"))
    # Sorted, merged exclusion intervals (avoid double counting).
    exclusions.sort()
    merged=[]
    for a,b in exclusions:
        if merged and a<=merged[-1][1]+1e-9:merged[-1]=(merged[-1][0],max(b,merged[-1][1]))
        else:merged.append((a,b))
    def inside(sta):return any(a-1e-8<=sta<=b+1e-8 for a,b in merged)
    def structure_at(sta):
        types={kind for a,b,kind in typed_exclusions if a-1e-8<=sta<=b+1e-8}
        return "/".join(kind for kind in ("교량","터널") if kind in types)
    def delta(sta):return alignment.evaluate(sta)[0]-ground.at(sta)[0]
    # Keep the design and ground breakpoints, curve boundaries, 25m stations,
    # and all structure boundaries to avoid smearing across them.
    stations=[start,end]
    stations += [s for s in ground.xs if start<s<end]
    stations += [s for s,_ in alignment.vips if start<s<end]
    for curve in alignment.curves:
        stations += [s for s in (curve.pvc_sta,curve.pvt_sta) if start<s<end]
    stations += [s for pair in merged for s in pair if start<s<end]
    k=math.ceil(start/interval)
    while k*interval<end:
        stations.append(k*interval);k+=1
    stations=sorted(set(round(s,9) for s in stations))
    # Detect cut/fill zero crossings even between regular station boundaries.
    extra=[]
    for a,b in zip(stations,stations[1:]):
        da,db=delta(a),delta(b)
        if da*db<0:
            lo,hi=a,b
            for _ in range(48):
                mid=(lo+hi)/2
                if delta(lo)*delta(mid)<=0:hi=mid
                else:lo=mid
            extra.append((lo+hi)/2)
    stations=sorted(set(stations+extra))
    sections=[]
    for sta in stations:
        z=ground.at(sta)[0]
        p=alignment.evaluate(sta)[0]
        sections.append(section_at(sta,z,p,width,cut_slope,fill_slope,inside(sta),structure_at(sta)))
    intervals=[]
    for left,right in zip(sections,sections[1:]):
        a,b=left.sta,right.sta
        middle=(a+b)/2
        excluded=inside(middle)
        cut=fill=0.
        if not excluded:
            # Trapezoidal average-end-area volume on each sign-consistent subinterval.
            v=(left.area_m2+right.area_m2)*(b-a)/2
            if delta(middle)<0:cut=v
            elif delta(middle)>0:fill=v
        intervals.append(EarthworkInterval(a,b,cut,fill,excluded))
    return EarthworkResult(tuple(sections),tuple(intervals),
        sum(v.cut_m3 for v in intervals),sum(v.fill_m3 for v in intervals),
        sum(v.end_sta-v.start_sta for v in intervals if v.excluded),
        width,cut_slope,fill_slope)


def export_earthwork_xlsx(result,path):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    wb=Workbook()
    ws=wb.active;ws.title='측점별 가상횡단'
    ws.append(['STA(m)','지반고(m)','계획고(m)','성토고(+)/절토고(-)(m)',
               '구분','단면적(m²)','구조물 경계/구간'])
    for s in result.sections:
        ws.append([s.sta,s.ground_el,s.plan_el,s.height,s.kind,s.area_m2,
                   s.structure_type])
    v=wb.create_sheet('구간별 토공량')
    v.append(['시점 STA(m)','종점 STA(m)','구간길이(m)','절토량(m³)','성토량(m³)','구조물 제외'])
    for r in result.intervals:
        v.append([r.start_sta,r.end_sta,r.end_sta-r.start_sta,r.cut_m3,r.fill_m3,
                  '제외' if r.excluded else ''])
    total=wb.create_sheet('총괄표',0)
    for row in [ ['항목','값','단위/비고'],['산출방법','SIMPLE','수평 지반 가정 / 개략 추정'],
                 ['노반 폭',result.roadbed_width,'m'],['절토 사면',result.cut_slope,'수평:수직 = m:1'],
                 ['성토 사면',result.fill_slope,'수평:수직 = m:1'],
                 ['절토량',result.cut_m3,'m³'],['성토량',result.fill_m3,'m³'],
                 ['구조물 제외 연장',result.excluded_length_m,'m'],
                 ['주의','실제 횡단 지형 미반영','편절·편성토 및 지형 변동 오차 가능'],
                 ['구조물','교량/터널 본체 제외','갱구·접속부 별도 토공 미산출']]:total.append(row)
    for sheet in wb:
        sheet.freeze_panes='A2'
        sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:
            cell.fill=PatternFill('solid',fgColor='17365D')
            cell.font=Font(bold=True,color='FFFFFF')
            cell.alignment=Alignment(vertical='center')
        for col in sheet.columns:
            letter=col[0].column_letter
            sheet.column_dimensions[letter].width=min(42,max(17,max(len(str(c.value or '')) for c in col[:100])+3))
        if sheet != total:
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value,float):cell.number_format='#,##0.000'
    wb.save(path)
