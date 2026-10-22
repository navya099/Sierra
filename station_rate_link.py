"""U11: read-only station segment -> III-81 price recommendations.
No automatic selection for mixed or tunnel stations. No changes to U9 cost ledger.
"""
from station81_rough import RATES, eligible

def recommend(station, track_type):
    kinds={p['kind'] for p in station['pieces']}
    if kinds=={'EARTH'}: structure='토공부';status='추천 가능'
    elif kinds=={'BRIDGE'}: structure='교량부';status='추천 가능'
    else: structure=None;status='사용자 검토 (복합 또는 터널 정거장)'
    rates=[r for r in RATES if eligible(r,track_type) and (structure is None or r['structure']==structure)] if structure else []
    return {'station':station['name'],'sequence':station['sequence'],'classification':station['classification'],
            'status':status,'structure':structure,'rates':rates}

def validate_selection(station, code, track_type):
    rec=recommend(station,track_type)
    if rec['status']!='추천 가능':
        raise ValueError(f"{station['name']}: 복합/터널 정거장은 표 Ⅲ-81 단가를 자동 확정할 수 없습니다.")
    row=next((r for r in rec['rates'] if r['code']==code),None)
    if row is None:raise ValueError(f"{station['name']}: 구조형식 또는 선로형식에 맞지 않는 단가입니다.")
    return row

def add_link_sheet(wb, entries):
    from openpyxl.styles import Font,PatternFill
    ws=wb.create_sheet('정거장단가연계')
    ws.append(['정거장명','구조물 순서','판정','단가 추천 상태','사용자 선택 단가코드','적용조건','수량(개소)','단가(원/개소)','금액(원)','비고'])
    total=0
    for station, rec, row in entries:
        amount=row['price_krw'] if row else 0;total+=amount
        ws.append([station['name'],station['sequence'],station['classification'],rec['status'],row['code'] if row else '',
                   (f"{row['structure']} / {row['service']} / {row['platform']} / {('일반부' if row['region']=='RURAL' else '도시부')}") if row else '',
                   1 if row else 0,amount if row else None,amount if row else None,
                   'KDI 2019 표 Ⅲ-81 / 사용자 확정 / 본선 공사비 미차감' if row else '미산정: 구조형식 검토 또는 단가 미선택'])
    ws.append(['합계','','','','','','','',total,'U9 노반정거장개략 시트와 별도 참고 산정 · 중복 합산 금지'])
    ws.freeze_panes='A2'
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
    for col,w in {'A':20,'B':26,'C':26,'D':34,'E':42,'F':65,'G':17,'H':24,'I':24,'J':85}.items():ws.column_dimensions[col].width=w
    for cells in ws.iter_rows(min_row=2):
        for c in cells[7:9]:c.number_format='#,##0'
    return ws
