"""v12.6.1 preliminary bridge/tunnel inventory. No earthwork or material volume."""
from dataclasses import dataclass
import math
from bridge_girder import span_ranges

@dataclass(frozen=True)
class StructureQuantity:
    kind: str
    name: str
    start_sta: float
    end_sta: float
    length_m: float
    span_count: int
    span_length_m: float
    abutment_count: int
    pier_count: int
    girder_count: int | None
    source: str

def calculate_structure_quantities(segments, girder_settings, girders_per_span=None):
    """One row per structure, no sheet-split double counting.

    Girder count requires an explicitly supplied, validated cross-section
    quantity; the longitudinal girder polygons represent spans, NOT a count
    of transverse girders.
    """
    girder_settings.validate()
    if girders_per_span is not None and (type(girders_per_span) is not int or girders_per_span < 1):
        raise ValueError('경간당 거더 개수는 1 이상의 정수여야 합니다.')
    rows=[]
    for seg in segments:
        if seg.kind not in ('BRIDGE','TUNNEL'): raise ValueError(f'알 수 없는 구조물 종류: {seg.kind}')
        a,b=seg.start_sta,seg.end_sta
        if not all(math.isfinite(v) for v in (a,b)) or b<=a:
            raise ValueError(f'{seg.name}: 구조물 측점 범위가 잘못되었습니다.')
        if seg.kind=='BRIDGE':
            spans=span_ranges(a,b,girder_settings.girder_length)
            n=len(spans)
            rows.append(StructureQuantity(seg.kind,seg.name,a,b,b-a,n,(b-a)/n,
                                          2,n-1,n*girders_per_span if girders_per_span else None,seg.source))
        else:
            rows.append(StructureQuantity(seg.kind,seg.name,a,b,b-a,0,0,0,0,None,seg.source))
    return rows

def export_quantities_xlsx(rows, path):
    """Export basic structure inventory using widely available openpyxl."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    detail = wb.active
    detail.title = '구조물별 수량'
    summary = wb.create_sheet('집계')
    headers = ['구분', '명칭', '시점 STA(m)', '종점 STA(m)', '연장(m)',
               '경간수(EA)', '경간당 연장(m)', '교대(EA)', '교각(EA)',
               '거더(EA)', '생성 출처']
    detail.append(headers)
    for r in rows:
        detail.append([
            '교량' if r.kind == 'BRIDGE' else '터널', r.name,
            r.start_sta, r.end_sta, r.length_m,
            r.span_count or None, r.span_length_m or None,
            r.abutment_count or None, r.pier_count or None,
            r.girder_count, r.source,
        ])
    for row in detail.iter_rows(min_row=2):
        for idx in (2, 3, 4, 6):
            row[idx].number_format = '#,##0.000'
    detail.freeze_panes = 'A2'
    detail.auto_filter.ref = detail.dimensions
    for i in range(1, 12):
        detail.column_dimensions[get_column_letter(i)].width = 22 if i == 2 else 17

    summary.append(['항목', '교량', '터널'])
    groups = {kind: [r for r in rows if r.kind == kind] for kind in ('BRIDGE', 'TUNNEL')}
    metrics = [
        ('개소(EA)', lambda r: 1),
        ('총연장(m)', lambda r: r.length_m),
        ('경간수(EA)', lambda r: r.span_count),
        ('교대수(EA)', lambda r: r.abutment_count),
        ('교각수(EA)', lambda r: r.pier_count),
    ]
    for label, getter in metrics:
        summary.append([label] + [sum(getter(r) for r in groups[k]) for k in ('BRIDGE', 'TUNNEL')])
    summary['B3'].number_format = '#,##0.000'
    summary['C3'].number_format = '#,##0.000'
    summary['A8'], summary['B8'], summary['C8'] = (
        '산출 기준', '실제 구조물 목록 기준', '공사비·재료 물량 제외')
    summary['A9'], summary['B9'], summary['C9'] = (
        '주의', '교대 2개/교량, 교각=경간수-1', '거더 개수는 별도 입력 시만 산출')
    for col in 'ABC':
        summary.column_dimensions[col].width = 37

    for ws in (detail, summary):
        for cell in ws[1]:
            cell.fill = PatternFill(fill_type='solid', fgColor='17365D')
            cell.font = Font(bold=True, color='FFFFFF')
            cell.alignment = Alignment(vertical='center')
        ws.row_dimensions[1].height = 25
    wb.save(str(path))
