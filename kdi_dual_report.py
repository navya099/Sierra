"""U16-R3: one normalized cost ledger; XLSX review summary and submission PDF."""
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

SECTIONS = ['노반 본선','노반 정거장','궤도','전차선','전력','송전선로','변전설비','신호','통신','건축','기계설비']

def _d(value):
    d=Decimal(str(value))
    if not d.is_finite() or d<0: raise ValueError('공사비 금액/수량 오류')
    return d

def build_ledger(road_rows, station_records, system_records):
    ledger=[]
    for r in road_rows:
        ledger.append(dict(section='노반 본선',sub=str(r.group),name=str(r.description),code=str(r.code),unit=str(r.unit),quantity=_d(r.quantity),unit_price=_d(r.unit_price),amount=_d(r.amount)))
    for r in station_records:
        ledger.append(dict(section='노반 정거장',sub='정거장',name=str(r['name']),code=str(r['code']),unit='개소',quantity=_d(r['quantity']),unit_price=_d(r['unit_price']),amount=_d(r['amount'])))
    for r in system_records:
        ledger.append(dict(section=str(r['group']),sub='',name=str(r['name']),code=str(r['code']),unit=str(r['unit']),quantity=_d(r['quantity']),unit_price=_d(r['unit_price']),amount=_d(r['amount'])))
    return ledger

def add_review_sheet(wb, ledger):
    from openpyxl.styles import Font, PatternFill, Alignment
    name='통합공사비집계'
    if name in wb: del wb[name]
    ws=wb.create_sheet(name,0)
    ws.append(['통합 개략공사비 집계 · KDI 2019년 기준 · 산정분만 표시'])
    ws.merge_cells('A1:G1')
    ws.append(['공종','세부공종','산정항목 / 단가코드','단위','수량','단가(원)','금액(원)'])
    totals=defaultdict(lambda:Decimal(0))
    for r in ledger:
        ws.append([r['section'],r['sub'],r['name']+' / '+r['code'],r['unit'],float(r['quantity']),float(r['unit_price']),float(r['amount'])])
        totals[r['section']]+=r['amount']
    ws.append(['공종별 소계'])
    for sec in SECTIONS+sorted(set(totals)-set(SECTIONS)):
        if sec in totals: ws.append([sec,'소계',None,None,None,None,float(totals[sec])])
    grand=sum(totals.values(),Decimal(0))
    ws.append(['전체 산정분 합계',None,None,None,None,None,float(grand)])
    ws.append(['※ 미산정 공종은 집계에서 제외하며, 실제 산정 0원은 그대로 표시합니다. 전체 사업비가 아닌 산정분 합계입니다.'])
    ws.merge_cells(start_row=ws.max_row,start_column=1,end_row=ws.max_row,end_column=7)
    for row in ws:
        for c in row:
            c.font=Font(name='맑은 고딕',size=10,bold=(c.row<=2 or c.column==1 and ('소계' in str(ws.cell(c.row,2).value) or '합계' in str(c.value))))
            if c.row<=2: c.fill=PatternFill('solid',fgColor='DCE8F5')
            if c.column in (5,6,7):c.number_format='#,##0.######'
            c.alignment=Alignment(vertical='center')
    for col,width in {'A':20,'B':18,'C':64,'D':12,'E':14,'F':20,'G':22}.items():ws.column_dimensions[col].width=width
    ws.freeze_panes='D3';ws.sheet_view.showGridLines=False
    ws.auto_filter.ref=f'A2:G{2+len(ledger)}'
    return grand

def export_pdf(path, ledger, *, line_km=None, station_count=None):
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, KeepTogether
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from xml.sax.saxutils import escape
    import os
    fonts=[r'C:\Windows\Fonts\malgun.ttf',r'C:\Windows\Fonts\gulim.ttc', '/usr/share/fonts/truetype/nanum/NanumGothic.ttf']
    font=next((p for p in fonts if os.path.isfile(p)),None)
    if not font:raise RuntimeError('PDF 출력용 한글 글꼴을 찾을 수 없습니다 (맑은 고딕 또는 나눔고딕 필요).')
    pdfmetrics.registerFont(TTFont('KDIReportKorean',font))
    normal=ParagraphStyle('normalKDI',fontName='KDIReportKorean',fontSize=7.3,leading=10)
    small=ParagraphStyle('smallKDI',parent=normal,fontSize=6.4,leading=8.5)
    title=ParagraphStyle('titleKDI',parent=normal,fontSize=14,leading=20,alignment=1,spaceAfter=12)
    doc=SimpleDocTemplate(str(path),pagesize=landscape(A4),leftMargin=30,rightMargin=30,topMargin=35,bottomMargin=35,title='KDI 표 Ⅲ-117 공사비 내역')
    P=lambda s,sty=small:Paragraph(escape(str(s)),sty)
    story=[P('표 Ⅲ-117  총사업비 내역 (일반철도: 개략적 기준)',title)]
    overview=f'가격 기준: 2019년  |  금액 단위: 백만원  |  공사비 산정분 (미산정 제외)'
    if line_km is not None:overview+=f'  |  총연장: {line_km:.3f} km'
    if station_count is not None:overview+=f'  |  정거장: {station_count}개소'
    story.extend([P(overview,normal),Spacer(1,10)])
    headers=['공종','규격 / 산정항목','단위','수량','단가(백만원)','금액(백만원)','비고']
    data=[[P(h,normal) for h in headers]]
    style=[('BACKGROUND',(0,0),(-1,0),colors.HexColor('#D9E3ED')),('FONTNAME',(0,0),(-1,-1),'KDIReportKorean'),('GRID',(0,0),(-1,-1),0.35,colors.HexColor('#8293A2')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4),('ALIGN',(2,1),(5,-1),'RIGHT')]
    by=defaultdict(list)
    for r in ledger:by[r['section']].append(r)
    total=Decimal(0)
    for sec in SECTIONS+sorted(set(by)-set(SECTIONS)):
        entries=by.get(sec,[])
        idx=len(data);data.append([P(sec,normal),'','','','','',P('미산정' if not entries else '',small)])
        style.extend([('BACKGROUND',(0,idx),(-1,idx),colors.HexColor('#EAF0F5')),('SPAN',(0,idx),(5,idx))])
        for r in entries:
            data.append(['',P((r['sub']+' · ' if r['sub'] else '')+r['name']),P(r['unit']),P(f"{r['quantity']:,.4f}".rstrip('0').rstrip('.')),P(f"{r['unit_price']/1000000:,.6f}".rstrip('0').rstrip('.')),P(f"{r['amount']/1000000:,.6f}".rstrip('0').rstrip('.')),P(r['code'])])
        subtotal=sum((r['amount'] for r in entries),Decimal(0));total+=subtotal
        idx=len(data);data.append([P(sec+' 소계',normal),'','','','',P(f'{subtotal/1000000:,.6f}'.rstrip('0').rstrip('.') if entries else ''),''])
        style.extend([('BACKGROUND',(0,idx),(-1,idx),colors.HexColor('#F4F6F8')),('SPAN',(0,idx),(4,idx))])
    idx=len(data);data.append([P('공사비 산정분 합계',normal),'','','','',P(f'{total/1000000:,.6f}'.rstrip('0').rstrip('.')),P('미산정 제외')]);style.extend([('BACKGROUND',(0,idx),(-1,idx),colors.HexColor('#D9E3ED')),('SPAN',(0,idx),(4,idx))])
    table=Table(data,colWidths=[95,216,43,62,88,92,179],repeatRows=1,hAlign='CENTER',splitByRow=1)
    table.setStyle(TableStyle(style))
    story.extend([table,Spacer(1,8),P('주: 실제 산정 0원은 0으로 표시하며, 미산정 항목은 공란으로 표시합니다. 전체 사업비가 아닌 산정된 공사비만 포함합니다.',normal)])
    def footer(canvas,doc):
        canvas.setFont('KDIReportKorean',8)
        canvas.drawRightString(landscape(A4)[0]-30,19,f'{doc.page} 쪽')
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return total
