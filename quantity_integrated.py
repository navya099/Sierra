"""v12.6.5 integrated quantity workbook; reuse v12.6.4 calculation results."""
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

def _tab(wb,name,header,rows):
    ws=wb.create_sheet(name);ws.append(header)
    for row in rows:ws.append(list(row))
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:
        cell.fill=PatternFill('solid',fgColor='17365D')
        cell.font=Font(color='FFFFFF',bold=True)
        cell.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=34
    for col in range(1,len(header)+1):
        values=[str(ws.cell(i,col).value or '') for i in range(1,min(ws.max_row,80)+1)]
        ws.column_dimensions[get_column_letter(col)].width=min(48,max(15,max(map(len,values))+3))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value,float):c.number_format='#,##0.000'
    return ws

def export_integrated_xlsx(path, earth, structures, bridges, tunnels,
                           bridge_section_path=None,tunnel_section_path=None):
    """Missing DXF-derived volumes remain blank; never silently become zero."""
    wb=Workbook();wb.remove(wb.active)
    b=list(bridges);t=list(tunnels);s=list(structures)
    def optional_sum(seq,attr):
        vals=[getattr(x,attr) for x in seq]
        return sum(vals) if vals and all(v is not None for v in vals) else None
    summary=[
        ('토공','절토량',earth.cut_m3,'m³','개략 / 수평지반 가정'),
        ('토공','성토량',earth.fill_m3,'m³','개략 / 수평지반 가정'),
        ('토공','구조물 제외 연장',earth.excluded_length_m,'m','교량·터널 본체 제외'),
        ('교량','개소',len(b),'EA',''),('교량','총연장',sum(x.length_m for x in b),'m',''),
        ('교량','경간',sum(x.spans for x in b),'EA',''),
        ('교량','교대',sum(x.abutments for x in b),'EA',''),
        ('교량','교각',sum(x.piers for x in b),'EA',''),
        ('교량','거더 개수',optional_sum(b,'girder_lines'),'EA','단면 DXF 없으면 미산출'),
        ('교량','슬래브 체적',optional_sum(b,'slab_volume_m3'),'m³','단면 DXF 없으면 미산출'),
        ('교량','거더 체적',optional_sum(b,'girder_volume_m3'),'m³','단면 DXF 없으면 미산출'),
        ('교량','코핑 개략 체적',optional_sum(b,'cap_volume_m3'),'m³','직육면체 가정'),
        ('교량','기초 개략 체적',optional_sum(b,'footing_volume_m3'),'m³','직육면체 가정'),
        ('터널','개소',sum(x.kind=='TUNNEL' for x in s),'EA',''),('터널','총연장',sum(x.length_m for x in s if x.kind=='TUNNEL'),'m',''),
        ('터널','굴착 체적',optional_sum(t,'excavation_m3'),'m³','사용자 단면 DXF 기준'),
        ('터널','라이닝 체적',optional_sum(t,'lining_m3'),'m³','사용자 단면 DXF 기준'),
        ('터널','내공 체적',optional_sum(t,'inner_space_m3'),'m³','참고용 내공 공간'),
    ]
    _tab(wb,'수량총괄표',['공종','항목','수량','단위','산출기준'],summary)
    _tab(wb,'토공수량',['항목','수량','단위','산출방법'],[
        ('절토량',earth.cut_m3,'m³','평균단면법'),('성토량',earth.fill_m3,'m³','평균단면법'),
        ('구조물 제외 연장',earth.excluded_length_m,'m','교량/터널')])
    _tab(wb,'토공상세',['시점 STA(m)','종점 STA(m)','구간길이(m)','절토량(m³)','성토량(m³)','구조물 제외'],[
        (x.start_sta,x.end_sta,x.end_sta-x.start_sta,x.cut_m3,x.fill_m3,'제외' if x.excluded else '') for x in earth.intervals])
    _tab(wb,'측점별 가상횡단',['STA(m)','지반고(m)','계획고(m)','성토고(+)/절토고(-)(m)','구분','단면적(m²)','구조물 유형'],[
        (x.sta,x.ground_el,x.plan_el,x.height,x.kind,x.area_m2,x.structure_type) for x in earth.sections])
    _tab(wb,'교량수량',['교량명','시점 STA(m)','종점 STA(m)','연장(m)','경간(EA)','교대(EA)','교각(EA)',
        '거더 개수(EA)','슬래브 단면적(m²)','거더 합계 단면적(m²)','슬래브 체적(m³)','거더 체적(m³)',
        '코핑 개략체적(m³)','기초 개략체적(m³)','출처'],[
        (x.name,x.start_sta,x.end_sta,x.length_m,x.spans,x.abutments,x.piers,x.girder_lines,
         x.slab_area_m2,x.girder_area_m2,x.slab_volume_m3,x.girder_volume_m3,x.cap_volume_m3,x.footing_volume_m3,x.source) for x in b])
    _tab(wb,'터널수량',['터널명','시점 STA(m)','종점 STA(m)','연장(m)','굴착 단면적(m²)','내공 단면적(m²)',
        '라이닝 단면적(m²)','굴착 체적(m³)','라이닝 체적(m³)','내공 체적(m³)','출처'],[
        (x.name,x.start_sta,x.end_sta,x.length_m,x.outer_area_m2,x.inner_area_m2,x.lining_area_m2,
         x.excavation_m3,x.lining_m3,x.inner_space_m3,x.source) for x in t])
    _tab(wb,'구조물현황',['구분','명칭','시점 STA(m)','종점 STA(m)','연장(m)','경간(EA)','경간당 연장(m)',
        '교대(EA)','교각(EA)','거더(EA)','출처'],[
        ('교량' if x.kind=='BRIDGE' else '터널',x.name,x.start_sta,x.end_sta,x.length_m,
         x.span_count or None,x.span_length_m or None,x.abutment_count or None,x.pier_count or None,
         x.girder_count,x.source) for x in s])
    _tab(wb,'산출기준',['구분','항목','적용값 / 기준','주의'],[
        ('토공','산출방법','SIMPLE / 수평 지반 가정','실제 횡단 물량이 아닌 개략 수량'),
        ('토공','노반 폭(m)',earth.roadbed_width,''),
        ('토공','절토 사면 (수평:수직)',earth.cut_slope,'m:1'),
        ('토공','성토 사면 (수평:수직)',earth.fill_slope,'m:1'),
        ('토공','구조물 구간','교량/터널 본체 제외','갱구·접속부 별도 물량 미산출'),
        ('교량','단면 DXF',bridge_section_path or '미지정','미지정 시 슬래브/거더 물량 공란'),
        ('교량','슬래브/거더','단면적 × 교량 연장','일정 단면 가정'),
        ('교량','코핑/기초','직육면체 가정','실제 3D 형상과 차이 가능'),
        ('교량','미산출','교대 콘크리트·교각 기둥·철근 등',''),
        ('터널','단면 DXF',tunnel_section_path or '미지정','미지정 시 터널 체적 미산출'),
        ('터널','굴착','외곽 단면적 × 터널 연장','DXF mm → m²'),
        ('터널','라이닝','(외곽-내공) 면적 × 터널 연장',''),
        ('터널','미산출','갱구부·여굴·부속시설',''),
        ('전체','집계 원칙','미산출은 공란','0으로 대체하지 않음'),
    ])
    wb.active=0
    wb.save(str(path))
