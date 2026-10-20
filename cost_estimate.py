"""v12.7.2 preliminary cost engine. Explicit rough selection; no invented engineering assumptions."""
from dataclasses import dataclass
from math import isfinite
from quantity_price import load_catalog, load_config, validate_quantity_price, effective_price

@dataclass(frozen=True)
class CostRow:
    group: str
    code: str
    description: str
    quantity: float
    unit: str
    unit_price: float
    amount: float
    basis: str


def available_rough(catalog=None):
    catalog=catalog or load_catalog()
    return {g:[x for x in catalog['items'] if x['group']==g and x['method']=='ROUGH' and x.get('source_price_krw') is not None] for g in ('토공','교량','터널')}


def estimate_rough(earth, structures, selections, catalog=None, config=None):
    """Select exactly one verified KDI rough price for each applicable group.

    Earthwork length comes from nonexcluded calculation intervals, not earthwork volume.
    Bridge/tunnel lengths are obtained from structure segments, with overlap rejection.
    No track count, foundation type, terrain type or NATM pattern is inferred.
    """
    catalog=catalog or load_catalog();config=config or load_config()
    lookup={x['code']:x for x in catalog['items']}
    segments=sorted([(float(s.start_sta),float(s.end_sta),s.kind) for s in structures if s.kind in ('BRIDGE','TUNNEL')])
    for a,b,k in segments:
        if not isfinite(a) or not isfinite(b) or b<=a:raise ValueError(f'{k}: 잘못된 구조물 연장')
    for (_,b,_),(a,_,_) in zip(segments,segments[1:]):
        if a < b-1e-8:raise ValueError('교량/터널 구조물 구간이 중첩되어 공사비 중복 가능성이 있습니다.')
    intervals=list(earth.intervals)
    earth_m=sum(float(x.end_sta-x.start_sta) for x in intervals if not x.excluded)
    lengths={'토공':earth_m/1000,'교량':sum(b-a for a,b,k in segments if k=='BRIDGE')/1000,'터널':sum(b-a for a,b,k in segments if k=='TUNNEL')/1000}
    if any(not isfinite(v) or v<0 for v in lengths.values()):raise ValueError('산출 연장을 확인하세요.')
    rows=[];review=[];used={}
    for group,quantity in lengths.items():
        code=selections.get(group,'')
        if quantity<=1e-9:
            if code:review.append((group,'연장 0km: 선택 단가 미적용'))
            continue
        if not code:
            review.append((group,f'{quantity:.6f}km 미산정: 단가 미선택'))
            continue
        item=lookup.get(code)
        if item is None or item['group']!=group or item['method']!='ROUGH':
            raise ValueError(f'{group}: 잘못된 개략 단가코드 {code}')
        if item['unit']!='km':raise ValueError(f'{code}: 개략 연장 단위 km 필요')
        amount=validate_quantity_price(quantity,'km',item,config,used)
        price=effective_price(item,config)
        rows.append(CostRow(group,code,f"{item['name']} / {item['spec']} / {item['condition']}",quantity,'km',price,amount,
                            f"KDI 2021 표 {item['source_table']} / 보고서 p.{item['source_page']} / {item['source_year']}년 기준"))
    return rows,review,lengths


def export_cost_xlsx(path,rows,review,lengths,config=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    config=config or load_config()
    wb=Workbook();wb.remove(wb.active)
    def add(name,header,records):
        ws=wb.create_sheet(name);ws.append(header)
        for r in records:ws.append(list(r))
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D');c.alignment=Alignment(wrap_text=True)
        for i in range(1,len(header)+1):ws.column_dimensions[get_column_letter(i)].width=30 if i!=2 else 42
        for rr in ws.iter_rows(min_row=2):
            for c in rr:
                if isinstance(c.value,(float,int)):c.number_format='#,##0.00'
        return ws
    add('공사비총괄',['공종','금액(원)','상태'],[(g,sum(r.amount for r in rows if r.group==g),'산정' if any(r.group==g for r in rows) else ('미산정' if lengths[g]>0 else '해당없음')) for g in lengths]+[('산정분 합계',sum(r.amount for r in rows),'미산정 공종 제외: 총사업비 아님')])
    add('공사비상세',['공종','단가코드','적용조건','수량','단위','적용단가(원/단위)','금액(원)','단가출처'],[(r.group,r.code,r.description,r.quantity,r.unit,r.unit_price,r.amount,r.basis) for r in rows])
    add('미산정검토',['공종','검토사항'],review or [('전체','미산정 공종 없음')])
    add('산정기준',['항목','값','주의'],[
        ('설계 기준연도',config.get('design_year'),'가격 자동 보정 아님'),('원문 기준연도',2019,'KDI 2021 일반철도'),
        ('기본 보정계수',config.get('default_factor',1),'1.0은 2019년 가격 미보정'),
        ('적용방식','개략 단가 × 구간연장(km)','상세 단가 및 부재별 수량과 중복 적용 금지'),
        ('토공 연장(km)',lengths['토공'],'구조물 제외 구간 / 수평지반 개략 토공량 엔진 범위'),
        ('교량 연장(km)',lengths['교량'],'구조물 목록'),('터널 연장(km)',lengths['터널'],'구조물 목록'),
        ('포함범위','선택된 KDI 개략 토공·교량·터널 공사비','용지비·차량비·부대비·제경비·VAT 등 별도 검토'),
        ('주의','공종별 조건은 사용자가 선택','단선/복선, 일반/도시, 기초형식 자동 추정 금지'),
        ('미산정','선택되지 않은 공종은 0원이 아님','총괄의 산정분 합계는 전체 사업비가 아님')])
    wb.active=0;wb.save(str(path))
