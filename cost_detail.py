"""v12.7.3-D1: explicit detailed KDI bill of quantities. No quantity guessing."""
from dataclasses import dataclass
from math import isfinite
from quantity_price import load_catalog,load_config,validate_quantity_price,effective_price

@dataclass(frozen=True)
class DetailRow:
    group:str
    code:str
    description:str
    quantity:float
    unit:str
    unit_price:float
    amount:float
    basis:str
    remark:str

def available_detail(catalog=None):
    catalog=catalog or load_catalog()
    return [x for x in catalog['items'] if x.get('method')=='DETAIL' and x.get('source_price_krw') is not None]

def estimate_detail(entries,catalog=None,config=None,rough_groups=()):
    """entries: [{code,quantity,remark}]. Missing quantities are not zero.

    No rough/detail combination within the same work group. Reuse price codes across different source quantities.
    """
    catalog=catalog or load_catalog();config=config or load_config()
    lookup={x['code']:x for x in available_detail(catalog)}
    rows=[];seen_sources=set();used={g:'ROUGH' for g in rough_groups}
    for n,entry in enumerate(entries,1):
        code=str(entry.get('code','')).strip()
        if not code:raise ValueError(f'{n}행: 단가코드 누락')
        source_id=str(entry.get('source_id') or '').strip()
        if source_id:
            if source_id in seen_sources:raise ValueError(f'{source_id}: 동일 수량 후보가 중복 등록되었습니다.')
            seen_sources.add(source_id)
        item=lookup.get(code)
        if item is None:raise ValueError(f'{code}: 검증된 세부 단가가 아닙니다.')
        q=float(entry['quantity'])
        if not isfinite(q) or q<=0:raise ValueError(f'{code}: 수량은 0보다 큰 유한수여야 합니다.')
        amount=validate_quantity_price(q,item['unit'],item,config,used)
        remark=str(entry.get('remark','')).strip()
        rows.append(DetailRow(item['group'],code,f"{item['name']} / {item['spec']} / {item['condition']}",q,item['unit'],effective_price(item,config),amount,
                    f"KDI 표 {item['source_table']} / 보고서 p.{item['source_page']} / {item['source_year']}년 기준",remark))
    return rows

def export_detail_xlsx(path,rows,config=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    config=config or load_config()
    wb=Workbook();wb.remove(wb.active)
    def sheet(name,header,records):
        ws=wb.create_sheet(name);ws.append(header)
        for record in records:ws.append(list(record))
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D');c.alignment=Alignment(wrap_text=True)
        for i in range(1,len(header)+1):ws.column_dimensions[get_column_letter(i)].width=24 if i>1 else 20
        return ws
    groups=('토공','교량','터널')
    sheet('세부공사비총괄',['공종','산정금액(원)','상태'],[(g,sum(r.amount for r in rows if r.group==g),'입력항목 산정' if any(r.group==g for r in rows) else '미산정') for g in groups]+[('산정분 합계',sum(r.amount for r in rows),'미입력 수량 제외, 총사업비 아님')])
    sheet('세부공사비내역',['공종','단가코드','KDI 항목·규격·조건','수량','단위','적용단가(원/단위)','금액(원)','원문 출처','수량 근거/비고'],[(r.group,r.code,r.description,r.quantity,r.unit,r.unit_price,r.amount,r.basis,r.remark) for r in rows])
    sheet('산정기준',['항목','내용'],[('설계연도',config.get('design_year')),('원문단가기준연도',2019),('기본보정계수',config.get('default_factor',1)),('수량 취득','사용자 명시 입력; 자동 추정하지 않음'),('중복방지','동일 공종 개략·세부 병용 금지'),('합계 주의','선택한 세부 항목만의 산정분 합계; 전체 공사비·사업비 아님'),('단가 주의','KDI 원문 포함범위 및 부가세 중복 검토 필요')])
    for ws in wb:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value,(int,float)):c.number_format='#,##0.00'
    wb.active=0;wb.save(str(path))
