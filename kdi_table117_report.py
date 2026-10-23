"""U16: KDI III-117 construction-only presentation sheet (2019 KRW -> million KRW).
Does not alter estimator inputs or underlying calculation sheets.
"""
from collections import defaultdict
from decimal import Decimal
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Source: KDI railway rough total-cost template III-117, pages 149-150.
SECTIONS = [
 ('노반 본선', [('토공',['일반구간','연약구간','기타']),('교량',['직접기초','말뚝기초','기타']),('터널',['본선','초장대 본선','경사갱','방재설비','기타']),('입체교차',['과선도로교','지하차도','보도육교','기타'])]),
 ('노반 정거장',[('토공',[]),('교량',[]),('지하',[])]),
 ('궤도',[]),('전차선',[]),('전력',[]),('송전선로',[]),('변전설비',[]),('신호',[]),('통신',[]),('건축',[]),('기계설비',[])
]

def _d(value):
 d=Decimal(str(value))
 if not d.is_finite() or d<0:raise ValueError('U16: 공사비 금액 오류')
 return d

def add_table117(wb, road_rows, station_records, system_records):
 """Adds presentation sheet; only amounts already calculated are included."""
 if 'KDI_III-117_공사비' in wb:del wb['KDI_III-117_공사비']
 ws=wb.create_sheet('KDI_III-117_공사비',0)
 ws.sheet_view.showGridLines=False
 ws.merge_cells('A1:G1');ws['A1']='표 Ⅲ-117 총사업비 내역 (일반철도: 개략적 기준) — 공사비'
 ws['A1'].font=Font(name='맑은 고딕',size=14,bold=True,color='FFFFFF')
 ws['A1'].fill=PatternFill('solid',fgColor='17365D')
 ws['A2']='가격기준';ws['B2']='2019년';ws['D2']='금액 단위';ws['E2']='백만원'
 ws.append(['공종','규격','단위','수량','단가(백만원)','금액(백만원)','비고'])
 for c in ws[3]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='366092');c.alignment=Alignment(horizontal='center')
 items=[]
 for r in road_rows:
  group=str(r.group)
  items.append(('노반 본선',group,str(r.description),r.unit,r.quantity,r.unit_price,r.amount,str(r.code)))
 for r in station_records:
  code=str(r['code']); typ='토공' if 'EARTH' in code else '교량' if 'BRIDGE' in code else '지하' if 'UNDER' in code or 'CUT' in code else '기타'
  items.append(('노반 정거장',typ,str(r['name']),'개소',r['quantity'],r['unit_price'],r['amount'],code))
 for r in system_records:
  group=str(r['group'])
  items.append((group,group,str(r['name']),str(r['unit']),r['quantity'],r['unit_price'],r['amount'],str(r['code'])))
 bysection=defaultdict(list)
 for item in items:bysection[item[0]].append(item)
 total=Decimal(0)
 def add_line(cells,kind='detail'):
  ws.append(cells);n=ws.max_row
  if kind!='detail':
   for c in ws[n]:c.font=Font(bold=True,color='17365D');c.fill=PatternFill('solid',fgColor='DCE6F1' if kind=='section' else 'E9EEF4')
  for col in (4,5,6):ws.cell(n,col).number_format='#,##0.######'
  ws.cell(n,7).alignment=Alignment(wrap_text=True,vertical='center')
  return n
 for section,subsections in SECTIONS:
  records=bysection.pop(section,[])
  add_line([section,None,None,None,None,None,'미산정 항목은 공란'], 'section')
  if subsections:
   known={name for name,_ in subsections}
   for sub,labels in subsections:
    subset=[r for r in records if r[1]==sub]
    add_line(['',sub,None,None,None,None,None],'sub')
    for r in subset:
     _,_,label,unit,qty,price,amt,code=r
     total+=_d(amt)
     add_line(['',label,unit,float(_d(qty)),float(_d(price)/1000000),float(_d(amt)/1000000),code])
   for r in records:
    if r[1] not in known:
     total+=_d(r[6]);add_line(['',r[2],r[3],float(_d(r[4])),float(_d(r[5])/1000000),float(_d(r[6])/1000000),r[7]])
  else:
   for r in records:
    total+=_d(r[6]);add_line(['',r[2],r[3],float(_d(r[4])),float(_d(r[5])/1000000),float(_d(r[6])/1000000),r[7]])
 # Unmapped categories must not be lost.
 for section,records in bysection.items():
  add_line([section,None,None,None,None,None,'원문 항목과 별도 검토 필요'],'section')
  for r in records:
   total+=_d(r[6]);add_line(['',r[2],r[3],float(_d(r[4])),float(_d(r[5])/1000000),float(_d(r[6])/1000000),r[7]])
 add_line(['공사비 산정분 합계',None,None,None,None,float(total/1000000),'전체 공사비 아님 · 미산정 공종 존재 가능'],'section')
 ws.append(['주의','원문 양식 참조: 표 Ⅲ-117','',None,None,None,'세부 규격별 원문 행 배치·단가 매핑은 추가 검증 필요. 산정값만 출력.'])
 for col,width in {'A':22,'B':45,'C':15,'D':16,'E':22,'F':24,'G':70}.items():ws.column_dimensions[col].width=width
 ws.freeze_panes='D4';ws.sheet_properties.pageSetUpPr.fitToPage=True
 ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.print_options.horizontalCentered=True
 ws.print_title_rows='1:3'
 return ws
