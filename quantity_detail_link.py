"""v12.7.3-D2: read-only integrated quantity workbook -> review candidates.
No inferred material, span type, foundation, or NATM pattern.
"""
from dataclasses import dataclass
from math import isfinite

@dataclass(frozen=True)
class Candidate:
    group: str
    item: str
    quantity: float
    unit: str
    source: str
    warning: str
    suggested_codes: tuple = ()


def read_candidates(path, catalog=None):
    from openpyxl import load_workbook
    from quantity_price import load_catalog
    catalog = catalog or load_catalog()
    wb = load_workbook(path, read_only=True, data_only=True)
    required = ('토공수량', '교량수량', '터널수량')
    missing = [s for s in required if s not in wb.sheetnames]
    if missing:
        wb.close()
        raise ValueError('v12.6.5 통합 수량산출서 형식이 아닙니다. 누락 시트: '+', '.join(missing))
    result=[]
    def rows(sheet):
        it=wb[sheet].iter_rows(values_only=True)
        header=next(it,None)
        if not header: raise ValueError(f'{sheet}: 빈 시트')
        return [dict(zip(header,r)) for r in it if any(x is not None for x in r)]
    def add(group,item,value,unit,source,warning):
        if value is None or value=='':return
        try:q=float(value)
        except (TypeError,ValueError):raise ValueError(f'{source}: 수량이 숫자가 아닙니다.')
        if not isfinite(q) or q<0:raise ValueError(f'{source}: 음수/무한 수량입니다.')
        if q==0:return
        codes=tuple(x['code'] for x in catalog['items'] if x.get('group')==group and x.get('method')=='DETAIL' and x.get('unit')==unit and x.get('source_price_krw') is not None)
        result.append(Candidate(group,item,q,unit,source,warning,codes))
    for r in rows('토공수량'):
        item=str(r.get('항목') or '')
        if item in ('절토량','성토량'):
            add('토공',item,r.get('수량'),str(r.get('단위') or ''),'토공수량/'+item,
                '가상 횡단 수량. 토질/공법 미분류이므로 단가 수동 확인 필요')
    for r in rows('교량수량'):
        name=str(r.get('교량명') or '이름 없음')
        add('교량',name+' 상부공',r.get('연장(m)'),'m','교량수량/'+name+'/연장',
            '교량 형식, 단·복선, 경간 길이 확인 후 상부공 단가 선택')
        for key,desc in [('교대(EA)','교대'),('교각(EA)','교각')]:
            # KDI substructure uses 개소, quantity sheet EA is count of locations.
            add('교량',name+' '+desc,r.get(key),'개소','교량수량/'+name+'/'+desc,
                '교량 경간, 단·복선, 직접/말뚝 기초 확인 후 하부공 단가 선택 (EA→개소는 개소 수 확인 필수)')
    for r in rows('터널수량'):
        name=str(r.get('터널명') or '이름 없음')
        add('터널',name+' NATM 미분류 전체연장',r.get('연장(m)'),'m','터널수량/'+name+'/연장',
            'NATM 패턴별 연장이 없음. 전체연장을 한 패턴으로 자동 확정하지 말 것; 패턴별 분할 필요')
    wb.close()
    return result


def review_selection(candidate,code,catalog=None):
    """Only permit explicitly selected matching-group and matching-unit detail prices."""
    from quantity_price import load_catalog
    catalog=catalog or load_catalog()
    items={x['code']:x for x in catalog['items']}
    item=items.get(code)
    if item is None or item.get('method')!='DETAIL' or item.get('source_price_krw') is None:
        raise ValueError('검증된 세부 단가를 선택해야 합니다.')
    if item['group']!=candidate.group or item['unit']!=candidate.unit:
        raise ValueError('공종 또는 단위가 일치하지 않습니다.')
    return {'code':code,'quantity':candidate.quantity,'remark':candidate.source+' / '+candidate.warning+' / 사용자 확인'}
